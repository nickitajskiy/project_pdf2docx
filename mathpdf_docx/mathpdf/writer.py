from io import BytesIO
from PIL import Image
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from .omml import latex_to_omml, normalize_latex
from .schema import Mode


def setup(doc):
    normal=doc.styles['Normal']
    normal.font.name='Times New Roman'; normal.font.size=Pt(12)
    lang=OxmlElement('w:lang'); lang.set(qn('w:val'),'ru-RU')
    normal.element.get_or_add_rPr().append(lang)
    normal.paragraph_format.space_after=Pt(4)
    normal.paragraph_format.line_spacing=1.15  # множитель, не фиксированная высота
    from docx.shared import RGBColor
    for name in ['Heading 1','Heading 2']:
        doc.styles[name].font.color.rgb=RGBColor(0,0,0)
        doc.styles[name].font.name='Times New Roman'
        doc.styles[name].font.size=Pt(12)
    sec=doc.sections[0]
    sec.page_width=Pt(595.276); sec.page_height=Pt(841.89)
    sec.left_margin=sec.right_margin=Pt(36)
    sec.top_margin=sec.bottom_margin=Pt(36)
    doc.core_properties.title='Конвертация математического PDF'
    doc.core_properties.author=''


def picture(paragraph, block, available=523):
    w,h=block.bbox[2]-block.bbox[0],block.bbox[3]-block.bbox[1]
    scale=min(1,available/max(w,1),700/max(h,1))
    run=paragraph.add_run()
    run.add_picture(block.crop,width=Pt(w*scale),height=Pt(h*scale))
    # Word по умолчанию ставит изображение на базовую линию. Небольшое смещение
    # вниз для формул с нижними индексами; высокие формулы оставляются отдельными.
    if h<27:
        prop=OxmlElement('w:position'); prop.set(qn('w:val'),str(-round(h*.22*2)))
        run._r.get_or_add_rPr().append(prop)
    for node in run._r.iter(qn('wp:docPr')):
        node.set('descr','Формула из PDF, не редактируется' if block.kind=='math' else 'Изображение из PDF')


def write_editable(paragraphs, output, mode, report, engine=None, pandoc=None):
    doc=Document(); setup(doc)
    for p in paragraphs:
        if p.page_number:
            for b in p.parts: b.method='page_number_separate'
            continue
        paragraph=doc.add_paragraph(style='Heading 2' if p.heading else 'Normal')
        paragraph.paragraph_format.left_indent=Pt(max(0,min(65,p.bbox[0]-42)))
        if p.heading:
            paragraph.paragraph_format.keep_with_next=True
            paragraph.paragraph_format.left_indent=Pt(0)
            paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if all(b.kind!='text' or not b.text.strip() for b in p.parts):
            paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for block in p.parts:
            if block.kind=='text':
                run=paragraph.add_run(block.text)
                run.font.name='Times New Roman'; run.font.size=Pt(max(6,min(24,block.size)))
                run.font.superscript=block.superscript; run.bold=block.bold
            elif block.kind=='math':
                if mode==Mode.OMML:
                    try:
                        with Image.open(block.crop) as im:
                            block.latex=engine.recognize(im.convert('RGB'))
                        block.latex_for_omml, block.latex_normalizations = normalize_latex(block.latex)
                        report.warn(block.page,'Распознанная формула требует сверки с оригиналом; синтаксис не доказывает правильность.')
                        paragraph._p.append(latex_to_omml(block.latex,pandoc))
                        block.method='omml_unverified'; report.omml+=1
                        continue
                    except Exception as e:
                        if block.latex and block.text_layer_latex:
                            try:
                                node=latex_to_omml(block.text_layer_latex,pandoc)
                            except Exception:
                                pass
                            else:
                                paragraph._p.append(node)
                                block.method='omml_text_layer_unverified'; report.omml+=1
                                block.latex_for_omml=block.text_layer_latex
                                msg='OCR отклонён; простая формула восстановлена из Unicode и геометрии текстового слоя PDF. Требуется сверка. Причина OCR: '+str(e)
                                block.warnings.append(msg); report.warn(block.page,msg)
                                continue
                        msg=f'Резервное изображение формулы: {type(e).__name__}: {e}'
                        block.warnings.append(msg); report.warn(block.page,msg)
                picture(paragraph,block); block.method='formula_image'; report.formula_images+=1
            else:
                picture(paragraph,block); block.method='image'
    # Номер исходной страницы сохранён в блоках отчёта, а нумерация DOCX собственная.
    footer=doc.sections[0].footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    doc.save(output)


def write_pages(images, sizes, output, report):
    doc=Document(); setup(doc)
    for i,(path,size) in enumerate(zip(images,sizes)):
        sec=doc.sections[0] if i==0 else doc.add_section(WD_SECTION_START.NEW_PAGE)
        sec.page_width=Pt(size[0]); sec.page_height=Pt(size[1])
        sec.top_margin=sec.bottom_margin=Pt(6); sec.left_margin=sec.right_margin=Pt(6)
        p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.line_spacing=1
        # Запас для обязательного абзаца Word, чтобы не появилась лишняя страница.
        w,h=size[0]-12,size[1]-20
        ratio=min(w/size[0],h/size[1])
        p.add_run().add_picture(str(path),width=Pt(size[0]*ratio),height=Pt(size[1]*ratio))
        report.page_images+=1
    doc.save(output)
