"""Ошибки из пользовательских скриншотов: края пределов и нативные дроби."""
from io import BytesIO
import json
import shutil
import pytest
from docx import Document
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
from mathpdf.converter import convert
from mathpdf.layout import prepare_regions, center_in
from mathpdf.omml import latex_to_omml, normalize_latex
from mathpdf.schema import Mode, Report
from test_converter import sample  # регистрирует доступный шрифт с кириллицей


@pytest.mark.parametrize('box,script', [
    ((78.6339,556.5102,134.1386,582.7572),
     {'text':'[','size':7.9701,'x0':72.653,'x1':75.005,'top':573.3211,'bottom':581.2912}),
    ((79.8905,666.9129,136.0386,693.4421),
     {'text':'[','size':7.9701,'x0':73.304,'x1':75.656,'top':682.9911,'bottom':690.9612}),
    ((77.0299,718.7834,135.5096,737.0818),
     {'text':'[','size':7.9701,'x0':72.003,'x1':74.355,'top':736.4571,'bottom':744.4272}),
])
def test_original_pdf_lower_limit_geometry(box, script):
    # Реальные координаты первых выпавших скобок страниц 3–4 образца.
    body=[{'text':'Т','size':12,'x0':400,'x1':409,'top':20+i*15,'bottom':32+i*15}
          for i in range(6)]
    regions=prepare_regions(body+[script],[box],595,842,Report('test'),1)
    assert len(regions)==1
    assert center_in(script,regions[0])
    assert regions[0][0]<=script['x0'] and regions[0][3]>=script['bottom']


@pytest.fixture
def lower_limit_pdf(sample,tmp_path):
    path=tmp_path/'Интеграл с областью.pdf'
    c=canvas.Canvas(str(path),pagesize=(595,842));c.setFont('TestRussian',12)
    c.drawString(45,780,'а)');c.drawString(80,780,'∫')
    c.drawString(95,780,'t dχ(t);')
    c.setFont('TestRussian',8);c.drawString(74,774,'[-1;1]')
    c.setFont('TestRussian',12);c.drawString(45,735,'Русский текст после формулы.')
    c.save();return path


class LimitEngine:
    def detect(self,image,size):return [(80,49,155,68)]
    def recognize(self,image):return r'\int_{[-1;1]}t\,d\chi(t)'


def test_limit_is_inside_image_not_a_separate_text_line(lower_limit_pdf):
    result=convert(lower_limit_pdf,engine=LimitEngine())
    doc=Document(BytesIO(result.docx))
    text=''.join(p.text for p in doc.paragraphs)
    assert '[' not in text and ']' not in text
    assert 'Русский текст после формулы.' in text
    assert len(doc.inline_shapes)==1


@pytest.mark.skipif(not shutil.which('pandoc'),reason='Pandoc отсутствует')
def test_limit_is_inside_native_word_integral(lower_limit_pdf):
    result=convert(lower_limit_pdf,Mode.OMML,engine=LimitEngine())
    doc=Document(BytesIO(result.docx))
    nary=list(doc.element.iter(qn('m:nary')))
    assert len(nary)==1 and len(doc.inline_shapes)==0
    sub=nary[0].find(qn('m:sub'))
    domain=''.join(t.text or '' for t in sub.iter(qn('m:t'))).replace('−','-')
    assert domain=='[-1;1]'
    assert not any(p.text.strip() in ['[',']','[0'] for p in doc.paragraphs)


@pytest.mark.parametrize('source,expected',[
    (r'\cfrac{1}{2+\cfrac{1}{2}}',r'\frac{1}{2+\frac{1}{2}}'),
    (r'\cfrac [l]{a}{b}',r'\frac{a}{b}'),
    (r'\dfrac{\partial f}{\partial t}',r'\frac{\partial f}{\partial t}'),
    (r'\tfrac12',r'\frac12'),
])
def test_fraction_style_compatibility_keeps_all_arguments(source,expected):
    actual,changes=normalize_latex(source)
    assert actual==expected and changes


def test_unknown_command_not_silently_rewritten():
    source=r'\cfractor{x}'
    assert normalize_latex(source)==(source,[])


@pytest.mark.skipif(not shutil.which('pandoc'),reason='Pandoc отсутствует')
def test_continued_fractions_are_native_and_nested():
    latex=(r'2,2+\frac12,2+\cfrac{1}{2+\cfrac12},'
           r'2+\cfrac{1}{2+\cfrac{1}{2+\cfrac12}},\ldots')
    math=latex_to_omml(latex)
    fractions=list(math.iter(qn('m:f')))
    assert len(fractions)==6
    assert any(list(f.find(qn('m:den')).iter(qn('m:f'))) for f in fractions)


@pytest.mark.skipif(not shutil.which('pandoc'),reason='Pandoc отсутствует')
def test_derivative_of_integral_is_native():
    latex=(r'\dfrac{d}{dt}\int_\Omega f(x,t)\,\mu(dx)='
           r'\int_\Omega\dfrac{\partial f(x,t)}{\partial t}\,\mu(dx)')
    math=latex_to_omml(latex)
    assert len(list(math.iter(qn('m:nary'))))==2
    assert len(list(math.iter(qn('m:f'))))==2
    assert 'Ω' in ''.join(t.text or '' for t in math.iter(qn('m:t')))


@pytest.mark.skipif(not shutil.which('pandoc'),reason='Pandoc отсутствует')
def test_raw_ocr_and_compatibility_version_are_both_in_report(lower_limit_pdf):
    class Fractions(LimitEngine):
        def recognize(self,image):return r'\cfrac{1}{2+\cfrac12}'
    result=convert(lower_limit_pdf,Mode.OMML,engine=Fractions())
    report=json.loads(result.report_json)
    block=next(b for b in report['blocks'] if b['kind']=='math')
    assert report['omml']==1 and report['formula_images']==0
    assert r'\cfrac' in block['latex'] and r'\cfrac' not in block['latex_for_omml']
    assert block['latex_normalizations'] and 'Для OMML:' in result.report_text


def test_owned_script_cannot_merge_neighboring_formulas():
    body=[{'text':'Т','size':12,'x0':400,'x1':409,'top':100+i*15,'bottom':112+i*15}
          for i in range(6)]
    script={'text':'1','size':8,'x0':75,'x1':79,'top':30,'bottom':38}
    boxes=[(70,10,110,28),(70,29,110,46)]
    regions=prepare_regions(body+[script],boxes,595,842,Report('test'),1)
    assert len(regions)==2
    assert all(b[3]-b[1]<22 for b in regions)


@pytest.mark.skipif(not shutil.which('pandoc'),reason='Pandoc отсутствует')
@pytest.mark.parametrize('latex,node',[(r'{\cal A}', 'm:r'),
                                      (r'\root n \of {nt}', 'm:rad')])
def test_other_equivalent_tex_syntax_is_native(latex,node):
    math=latex_to_omml(latex)
    assert list(math.iter(qn(node)))


def test_unsupported_ocr_letter_is_not_corrected_by_style_normalizer():
    # \l — не вариант оформления латинской l, нельзя исправлять его вслепую.
    assert normalize_latex(r'\l_p')==(r'\l_p',[])


@pytest.mark.skipif(not shutil.which('pandoc'),reason='Pandoc отсутствует')
def test_cases_trailing_empty_row_does_not_break_word_matrix():
    latex=r'F(x)=\begin{cases}0&x<0\\\sqrt{x}&x\geq0\\\end{cases}'
    math=latex_to_omml(latex)
    matrices=list(math.iter(qn('m:m')))
    assert len(matrices)==1
    rows=matrices[0].findall(qn('m:mr'))
    assert len(rows)==2 and all(len(row.findall(qn('m:e')))==2 for row in rows)
    assert list(math.iter(qn('m:rad')))


@pytest.mark.skipif(not shutil.which('pandoc'),reason='Pandoc отсутствует')
def test_single_aligned_integral_has_no_empty_rows_or_alignment_ampersand():
    latex=r'\begin{aligned} { } & { { } \int _ { [ 0 ; 1 ] } t ^ { 2 } \, d c ( t ) . } \\ \end{aligned}'
    math=latex_to_omml(latex)
    text=''.join(t.text or '' for t in math.iter(qn('m:t')))
    assert '&' not in text
    assert not list(math.iter(qn('m:eqArr')))
    assert len(list(math.iter(qn('m:nary'))))==1
    assert '[0;1]' in text


@pytest.mark.skipif(not shutil.which('pandoc'),reason='Pandoc отсутствует')
def test_empty_script_base_is_rejected_instead_of_word_placeholder():
    with pytest.raises(RuntimeError,match='индекс без основания'):
        latex_to_omml(r'{}^{c}')


@pytest.mark.skipif(not shutil.which('pandoc'),reason='Pandoc отсутствует')
@pytest.mark.parametrize('domain',[r'\Omega',r'\mathbb{R}',r'[-1;1]'])
def test_underset_integral_has_attached_body_and_no_placeholder(domain):
    math=latex_to_omml(r'\underset{'+domain+r'}{\int}\, f(x,t)\,\mu(dx)')
    operator=next(math.iter(qn('m:nary')))
    body=operator.find(qn('m:e'))
    assert 'f' in ''.join(body.itertext())
    assert operator.find(qn('m:naryPr')).find(qn('m:limLoc')).get(qn('m:val'))=='undOvr'


@pytest.mark.skipif(not shutil.which('pandoc'),reason='Pandoc отсутствует')
def test_empty_operator_body_is_not_accepted_as_complete_equation():
    with pytest.raises(RuntimeError,match='пустой шаблон'):
        latex_to_omml(r'\int_\Omega{}')
