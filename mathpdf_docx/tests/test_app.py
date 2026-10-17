"""Регрессии окна, фонового конвейера и сохранения результата."""
from io import BytesIO
import json
from pathlib import Path
from queue import Queue
import time
import tkinter as tk
import pytest
from docx import Document
from mathpdf.desktop import DesktopApp, conversion_worker, save_result
from mathpdf.schema import Mode
from test_converter import sample, FixtureEngine
from mathpdf.converter import convert


@pytest.fixture
def window():
    try:
        root = tk.Tk()
    except tk.TclError as exc:
        pytest.skip('Графическое окружение Tk недоступно: '+str(exc))
    root.withdraw()
    app = DesktopApp(root)
    yield app
    if app.busy:
        deadline=time.monotonic()+10
        while app.busy and time.monotonic()<deadline:
            root.update();time.sleep(.02)
    root.after_cancel(app.poll_id)
    root.destroy()


def wait_done(app):
    deadline=time.monotonic()+15
    while app.busy and time.monotonic()<deadline:
        app.root.update();time.sleep(.02)
    assert not app.busy


def test_window_modes_idle_no_models(window,monkeypatch):
    monkeypatch.setattr('mathpdf.math_engine.get_engine',lambda *a:pytest.fail('Окно не должно загружать модели'))
    assert len(window.combo['values'])==3
    assert not window.busy and window.result is None
    window.mode.set(Mode.PAGES.value);window.change_mode()
    assert 'не редактируются' in window.help['text']
    assert str(window.save_button['state'])=='disabled'
    window.root.update()


def test_worker_events_real_conversion(sample):
    events=Queue()
    conversion_worker(events,sample,Mode.IMAGES,50,100,
                      converter=lambda source,mode,**kw:convert(source,mode,engine=FixtureEngine(),**kw))
    records=[]
    while not events.empty():records.append(events.get())
    assert records[-1][0]=='done'
    assert any(kind=='progress' for kind,_ in records)
    assert len(Document(BytesIO(records[-1][1].docx)).inline_shapes)==2


def test_worker_error_is_reported(tmp_path):
    events=Queue();conversion_worker(events,tmp_path/'нет.pdf',Mode.PAGES,50,100)
    kind,message=events.get()
    # Проверка PDF может сначала сообщить этап обработки.
    while kind=='progress':kind,message=events.get()
    assert kind=='error' and message


def test_window_conversion_save_cancel_and_reports(window,sample,tmp_path,monkeypatch):
    window.source.set(str(sample));window.mode.set(Mode.PAGES.value)
    window.start();assert window.busy
    assert str(window.combo['state'])=='disabled'
    wait_done(window)
    assert window.result is not None and len(Document(BytesIO(window.result.docx)).inline_shapes)==1
    assert 'не редактируются' in window.report.get('1.0','end')
    assert str(window.save_button['state'])=='normal'
    monkeypatch.setattr('mathpdf.desktop.filedialog.asksaveasfilename',lambda **kw:'')
    window.save();assert window.result is not None
    target=tmp_path/'Результат с пробелами.docx'
    monkeypatch.setattr('mathpdf.desktop.filedialog.asksaveasfilename',lambda **kw:str(target))
    window.save()
    assert target.exists() and target.with_suffix('.report.txt').exists()
    assert json.loads(target.with_suffix('.report.json').read_text())['pages']==1
    assert len(Document(target).inline_shapes)==1


def test_window_failure_allows_retry(window,sample,tmp_path,monkeypatch):
    errors=[]
    monkeypatch.setattr('mathpdf.desktop.messagebox.showerror',lambda *a,**kw:errors.append(a))
    broken=tmp_path/'ошибка.pdf';broken.write_bytes(b'broken')
    window.source.set(str(broken));window.mode.set(Mode.PAGES.value)
    window.start();wait_done(window)
    assert errors and window.result is None and str(window.convert_button['state'])=='normal'
    window.source.set(str(sample));window.start();wait_done(window)
    assert window.result is not None


def test_save_source_protection(sample):
    result=convert(sample,Mode.PAGES)
    with pytest.raises(ValueError):save_result(result,sample,sample)
    with pytest.raises(ValueError):save_result(result,sample.with_suffix('.txt'),sample)


def test_close_during_processing_keeps_window(window,monkeypatch):
    notices=[]
    monkeypatch.setattr('mathpdf.desktop.messagebox.showinfo',lambda *a,**kw:notices.append(a))
    window.set_busy(True);window.close()
    assert notices and window.root.winfo_exists()
    window.set_busy(False)


def test_ide_launcher_independent_of_working_directory(tmp_path,monkeypatch):
    import importlib.util
    import os
    path=Path(__file__).parents[1]/'run.py'
    spec=importlib.util.spec_from_file_location('ide_launcher',path)
    launcher=importlib.util.module_from_spec(spec);spec.loader.exec_module(launcher)
    monkeypatch.chdir(tmp_path);monkeypatch.delenv('MATHPDF_MODEL_DIR',raising=False)
    calls=[]
    monkeypatch.setattr('mathpdf.desktop.main',lambda:calls.append(os.environ['MATHPDF_MODEL_DIR']))
    assert launcher.main()==0 and calls==[str(path.parent/'models')]


def test_changed_parameters_cannot_save_old_result(window,sample):
    window.source.set(str(sample));window.mode.set(Mode.PAGES.value)
    window.start();wait_done(window)
    assert window.result is not None
    window.mode.set(Mode.OMML.value);window.change_mode()
    assert window.result is None and str(window.save_button['state'])=='disabled'
    window.mode.set(Mode.PAGES.value);window.start();wait_done(window)
    assert window.result is not None
    window.source.set('другой.pdf')
    assert window.result is None and str(window.save_button['state'])=='disabled'
