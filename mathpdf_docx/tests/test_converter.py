from io import BytesIO
from pathlib import Path
import json
import shutil
from zipfile import ZipFile
import pytest
from docx import Document
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from mathpdf.converter import convert
from mathpdf.schema import Mode, ConversionError, Block, Paragraph, Report
from mathpdf.layout import build_paragraphs
from mathpdf.pdfio import inspect_pdf
from mathpdf.omml import latex_to_omml
from mathpdf.writer import write_editable

FONT_CANDIDATES=[Path('/System/Library/Fonts/Supplemental/Arial.ttf'),
                 Path('C:/Windows/Fonts/arial.ttf'),
                 Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')]

@pytest.fixture
def sample(tmp_path):
    font=next((p for p in FONT_CANDIDATES if p.exists()),None)
    if font is None: pytest.skip('Нужен установленный шрифт с кириллицей')
    pdfmetrics.registerFont(TTFont('TestRussian',str(font)))
    path=tmp_path/'Пример с пробелом.pdf'
    c=canvas.Canvas(str(path),pagesize=(595,842))
    c.setFont('TestRussian',12)
    c.drawString(45,780,'1'); c.setFont('TestRussian',8); c.drawString(52,785,'0')
    c.setFont('TestRussian',12); c.drawString(57,780,'. Докажите равенство')
    c.drawString(220,780,'x+1'); c.drawString(255,780,'для всех значений.')
    c.drawString(220,730,'x=2')
    c.drawString(45,700,'Русский текст после формулы.')
    c.drawString(292,20,'1')
    c.save(); return path

class FixtureEngine:
    """Известные области синтетического PDF. Не проверка точности Pix2Text."""
    def detect(self,image,size): return [(219,49,242,65),(219,99,245,116)]
    def recognize(self,image): return r'\frac{x+1}{2}'


def test_russian_superscript_inline_display_no_duplicates(sample):
    r=convert(sample,engine=FixtureEngine())
    d=Document(BytesIO(r.docx)); text='\n'.join(p.text for p in d.paragraphs)
    assert 'Докажите равенство' in text and 'Русский текст после формулы.' in text
    assert 'x+1' not in text and 'x=2' not in text
    assert len(d.inline_shapes)==2
    assert any(run.font.superscript and '0' in run.text for p in d.paragraphs for run in p.runs)
    report=json.loads(r.report_json)
    assert report['math_regions']==report['formula_images']==2
    assert any(b['method']=='page_number_separate' for b in report['blocks'])
    # Inline formula must share paragraph with editable text.
    assert any(p.text.startswith('1') and '<w:drawing>' in p._p.xml for p in d.paragraphs)


def test_pages_without_models(sample):
    r=convert(sample,Mode.PAGES)
    d=Document(BytesIO(r.docx)); assert len(d.inline_shapes)==1
    assert json.loads(r.report_json)['page_images']==1
    assert 'не редактируются' in r.report_text


def test_missing_model(sample,monkeypatch):
    def absent(*args): raise ConversionError('Модель недоступна')
    monkeypatch.setattr('mathpdf.converter.get_engine',absent)
    with pytest.raises(ConversionError,match='Модель'): convert(sample)
    assert convert(sample,Mode.PAGES).docx


def test_missing_pandoc_falls_back(sample,monkeypatch):
    monkeypatch.setattr('mathpdf.omml.shutil.which',lambda _:None)
    r=convert(sample,Mode.OMML,engine=FixtureEngine())
    report=json.loads(r.report_json)
    assert report['omml']==0 and report['formula_images']==2
    assert 'Pandoc не найден' in r.report_text


def test_bad_recognition_falls_back(sample):
    class Bad(FixtureEngine):
        def recognize(self,image): raise RuntimeError('сбой OCR')
    r=convert(sample,Mode.OMML,engine=Bad())
    assert json.loads(r.report_json)['formula_images']==2
    assert 'сбой OCR' in r.report_text

@pytest.mark.skipif(not shutil.which('pandoc'),reason='Pandoc отсутствует')
def test_real_pandoc_omml_structure(sample):
    r=convert(sample,Mode.OMML,engine=FixtureEngine())
    doc=Document(BytesIO(r.docx))
    assert len(list(doc.element.iter(qn('m:oMath'))))==2
    assert len(list(doc.element.iter(qn('m:f'))))==2
    report=json.loads(r.report_json)
    assert report['omml']==2 and report['formula_images']==0

@pytest.mark.skipif(not shutil.which('pandoc'),reason='Pandoc отсутствует')
def test_unsupported_latex_rejected():
    with pytest.raises(Exception): latex_to_omml(r'\notarealcommand{x}')

@pytest.mark.parametrize('payload',[b'',b'%PDF-1.7 broken',b'not a pdf'])
def test_invalid_pdf(tmp_path,payload):
    p=tmp_path/'bad.pdf'; p.write_bytes(payload)
    with pytest.raises(ConversionError): convert(p,Mode.PAGES)


def test_limits(sample):
    with pytest.raises(ConversionError): inspect_pdf(sample,max_pages=0)
    with pytest.raises(ConversionError): inspect_pdf(sample,max_mb=0)


def test_scan_requires_pages(tmp_path):
    p=tmp_path/'blank.pdf'; c=canvas.Canvas(str(p)); c.showPage(); c.save()
    with pytest.raises(ConversionError,match='текстового слоя'): convert(p,engine=FixtureEngine())
    assert convert(p,Mode.PAGES).docx


def test_encrypted(sample,tmp_path):
    from pypdf import PdfReader,PdfWriter
    w=PdfWriter(); w.append(PdfReader(sample)); w.encrypt('secret')
    p=tmp_path/'protected.pdf'
    with p.open('wb') as f: w.write(f)
    with pytest.raises(ConversionError): convert(p,Mode.PAGES)


def test_source_unchanged(sample):
    before=sample.read_bytes(); convert(sample,Mode.PAGES); assert sample.read_bytes()==before


def test_multiple_display_regions(sample):
    class Multi(FixtureEngine):
        def detect(self,image,size):
            return [(219,49,242,65),(219,99,245,116),(200,160,300,195),(180,210,320,260)]
    r=convert(sample,engine=Multi())
    assert json.loads(r.report_json)['formula_images']==4


def test_task_mark_not_detected_as_math(sample):
    class NumberDetector(FixtureEngine):
        def detect(self,image,size):
            return super().detect(image,size)+[(44,48,58,64)]
    r=convert(sample,engine=NumberDetector())
    doc=Document(BytesIO(r.docx))
    assert any(run.font.superscript and '0' in run.text for p in doc.paragraphs for run in p.runs)
    assert json.loads(r.report_json)['math_regions']==2


def test_combined_math_region_keeps_list_order(sample,tmp_path):
    path=tmp_path/'список.pdf'
    c=canvas.Canvas(str(path),pagesize=(595,842)); c.setFont('TestRussian',12)
    c.drawString(45,780,'а)'); c.drawString(75,780,'x=1')
    c.drawString(45,765,'б)'); c.drawString(75,765,'x=2'); c.save()
    class Combined:
        def detect(self,image,size): return [{'bbox':(74,49,100,82),'inline':True}]
    result=convert(path,engine=Combined())
    doc=Document(BytesIO(result.docx))
    ps=[p for p in doc.paragraphs if p.text.strip()]
    assert ps[0].text.startswith('а)') and ps[1].text.startswith('б)')
    assert len(doc.inline_shapes)==2
    assert '<w:drawing>' in ps[0]._p.xml and '<w:drawing>' in ps[1]._p.xml


def test_font_bbox_does_not_capture_previous_prose():
    from mathpdf.layout import prepare_regions
    c={'text':'(cid:112)','x0':100,'x1':110,'top':10,'bottom':50}
    report=Report('test')
    regions=prepare_regions([c],[(100,30,150,52)],595,842,report,1)
    assert len(regions)==1 and regions[0][1]>=29
    assert not report.warnings


def test_cli_never_overwrites_source(sample,monkeypatch):
    import sys
    from mathpdf.__main__ import main
    before=sample.read_bytes()
    monkeypatch.setattr(sys,'argv',['mathpdf','convert',str(sample),str(sample),'--mode','pages'])
    assert main()==1
    assert sample.read_bytes()==before


def test_columns_require_appearance_mode(sample,tmp_path):
    p=tmp_path/'две колонки.pdf'
    c=canvas.Canvas(str(p),pagesize=(595,842));c.setFont('TestRussian',12)
    for i in range(12):
        c.drawString(45,780-i*18,'Текст левой колонки')
        c.drawString(320,780-i*18,'Текст правой колонки')
    c.save()
    class NoMath:
        def detect(self,image,size):return []
    with pytest.raises(ConversionError,match='колонок'):convert(p,engine=NoMath())
    assert convert(p,Mode.PAGES).docx
